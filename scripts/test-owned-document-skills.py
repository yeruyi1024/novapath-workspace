"""Exercise StarPath helpers using the shipped document runtime, without network."""
import importlib.util
import os
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
sys.dont_write_bytecode = True


def module(name, relative):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


word = module('starpath_word', 'SKILLs/docx/scripts/word.py')
slides = module('starpath_slides', 'SKILLs/pptx/scripts/slides.py')
sheets = module('starpath_sheets', 'SKILLs/xlsx/scripts/sheets.py')
pdf = module('starpath_pdf', 'SKILLs/pdf/scripts/pdf_tools.py')


class Documents(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='starpath-owned-documents-')
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name) / '中文 folder'
        self.root.mkdir()

    def make_word(self):
        path = self.root / 'input.docx'
        document = Document()
        document.add_paragraph('Original text')
        document.add_table(rows=1, cols=1).cell(0, 0).text = '70'
        document.save(path)
        return path

    def make_deck(self):
        path = self.root / 'input.pptx'
        deck = Presentation()
        for title in ('First', 'Second'):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = title
            slide.notes_slide.notes_text_frame.text = f'Notes for {title}'
        deck.save(path)
        return path

    def make_pdf(self):
        path = self.root / 'input.pdf'
        document = canvas.Canvas(str(path))
        for title in ('First', 'Second'):
            document.drawString(72, 720, title)
            document.showPage()
        document.save()
        return path

    def test_word_comment_preserves_body_and_table(self):
        source = self.make_word()
        output = self.root / 'comment.docx'
        word.edit(source, output, 0, 'Check this', 'StarPath', False)
        result = word.inspect(output)
        self.assertEqual(result['paragraphs'], ['Original text'])
        self.assertEqual(result['tables'], [[['70']]])
        self.assertEqual(result['comments'][0]['text'], 'Check this')
        self.assertEqual(word.inspect(source)['comments'], [])

    def test_word_revision_preserves_unmodified_parts(self):
        source = self.make_word()
        output = self.root / 'revision.docx'
        word.edit(source, output, 0, 'Updated text', 'StarPath', True)
        result = word.inspect(output)
        self.assertEqual((result['insertions'], result['deletions']), (1, 1))
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
            for name in before.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(before.read(name), after.read(name))
            self.assertIn(b'Updated text', after.read('word/document.xml'))

    def test_word_rejects_rich_paragraph_without_writing(self):
        source = self.make_word()
        document = Document(source)
        document.paragraphs[0].add_run().add_break()
        document.save(source)
        output = self.root / 'revision.docx'
        with self.assertRaises(ValueError):
            word.edit(source, output, 0, 'Updated', 'StarPath', True)
        self.assertFalse(output.exists())

    def test_slide_replace_and_reorder_preserve_notes(self):
        source = self.make_deck()
        edited = self.root / 'edited.pptx'
        ordered = self.root / 'ordered.pptx'
        slides.replace(source, edited, 0, 0, 'Changed')
        slides.reorder(edited, ordered, [1, 0])
        result = slides.inspect(ordered)['slides']
        self.assertEqual([item['shapes'][0]['text'] for item in result], ['Second', 'Changed'])
        self.assertIn('Notes for First', result[1]['notes'])
        self.assertEqual(slides.inspect(source)['slides'][0]['shapes'][0]['text'], 'First')

    def test_slide_rejects_invalid_order(self):
        source = self.make_deck()
        output = self.root / 'bad.pptx'
        for order in ([], [-1], [2], [0, 0]):
            with self.assertRaises(ValueError):
                slides.reorder(source, output, order)
        self.assertFalse(output.exists())

    def test_spreadsheet_inspects_formulas_and_errors(self):
        source = self.root / 'input.xlsx'
        workbook = Workbook()
        workbook.active.append([70, '=A1*2', '#REF!'])
        workbook.save(source)
        result = sheets.inspect(source)['sheets'][0]
        self.assertEqual(result['formulas'], ['B1'])
        self.assertEqual(result['errors'], [{'cell': 'C1', 'value': '#REF!'}])

    @unittest.skipUnless(os.environ.get('STARPATH_LIBREOFFICE'), 'Local LibreOffice not configured')
    def test_spreadsheet_recalculates_actual_formula(self):
        source = self.root / 'input.xlsx'
        output = self.root / 'calculated.xlsx'
        workbook = Workbook()
        workbook.active.append([70, '=A1*2'])
        workbook.save(source)
        sheets.recalc(source, output)
        loaded = load_workbook(output, data_only=True)
        self.assertEqual(loaded.active['B1'].value, 140)
        loaded.close()

    def test_pdf_selection_and_merge(self):
        source = self.make_pdf()
        selected = self.root / 'selected.pdf'
        merged = self.root / 'merged.pdf'
        pdf.select(source, selected, '2')
        self.assertIn('Second', PdfReader(selected).pages[0].extract_text())
        pdf.merge([selected, source], merged)
        self.assertEqual(len(PdfReader(merged).pages), 3)

    def test_pdf_field_fill_and_unknown_field(self):
        source = self.root / 'form.pdf'
        output = self.root / 'filled.pdf'
        document = canvas.Canvas(str(source))
        document.acroForm.textfield(name='customer', x=72, y=600, width=200, height=24)
        document.showPage()
        document.save()
        pdf.fill(source, output, {'customer': 'StarPath'})
        self.assertEqual(PdfReader(output).get_fields()['customer']['/V'], 'StarPath')
        with self.assertRaises(ValueError):
            pdf.fill(source, self.root / 'bad.pdf', {'missing': 'value'})

    def test_outputs_never_overwrite_existing_files(self):
        source = self.make_pdf()
        before = source.read_bytes()
        with self.assertRaises(FileExistsError):
            pdf.select(source, source, '1')
        self.assertEqual(source.read_bytes(), before)

    def test_shared_xml_helper_rejects_external_entities(self):
        with self.assertRaises(ValueError):
            word.parse_xml(b'<!DOCTYPE a [<!ENTITY x SYSTEM "file:///private.txt">]><a>&x;</a>')


if __name__ == '__main__':
    unittest.main()
