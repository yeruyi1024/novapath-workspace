"""Exercise restored document entrypoints on temporary files, without model calls."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
SKILLS = ROOT / 'SKILLs'
sys.dont_write_bytecode = True


class RestoredDocuments(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='novapath-restored-documents-')
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name) / '中文 folder'
        self.root.mkdir()

    def run_tool(self, relative, *arguments):
        result = subprocess.run([sys.executable, str(SKILLS / relative), *map(str, arguments)],
                                text=True, capture_output=True, timeout=90)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        return result.stdout

    def make_word(self):
        source = self.root / 'input.docx'
        document = Document()
        paragraph = document.add_paragraph('Original ')
        paragraph.add_run('rich text').bold = True
        document.add_table(rows=1, cols=1).cell(0, 0).text = '70'
        document.save(source)
        return source

    def make_pdf(self, form=False):
        source = self.root / 'input.pdf'
        document = canvas.Canvas(str(source), pagesize=(600, 800))
        document.drawString(40, 750, 'Restored PDF document')
        if form:
            document.acroForm.textfield(name='name', x=100, y=600, width=200, height=25)
        document.showPage()
        document.save()
        return source

    def test_word_javascript_creation(self):
        output = self.root / 'javascript.docx'
        code = '''const fs = require('fs');
const {Document, Packer, Paragraph, TextRun} = require('docx');
const document = new Document({sections:[{children:[new Paragraph({children:[new TextRun({text:'Restored Word 中文',bold:true})]})]}]});
Packer.toBuffer(document).then(data=>fs.writeFileSync(process.argv[1],data)).catch(error=>{console.error(error);process.exitCode=1;});'''
        result = subprocess.run(['node', '-e', code, str(output)], cwd=SKILLS / 'docx',
                                text=True, capture_output=True, timeout=45)
        self.assertEqual(result.returncode, 0, result.stderr)
        paragraph = Document(output).paragraphs[0]
        self.assertEqual(paragraph.text, 'Restored Word 中文')
        self.assertTrue(paragraph.runs[0].bold)

    def test_word_unpack_schema_validation_and_repack(self):
        source = self.make_word()
        directory = self.root / 'unpacked'
        self.run_tool('docx/ooxml/scripts/unpack.py', source, directory)
        self.run_tool('docx/ooxml/scripts/validate.py', directory, '--original', source)
        output = self.root / 'packed.docx'
        self.run_tool('docx/ooxml/scripts/pack.py', directory, output)
        self.assertEqual(Document(output).paragraphs[0].text, 'Original rich text')
        self.assertEqual(Document(output).tables[0].cell(0, 0).text, '70')

    def test_word_original_editor_adds_comment_to_rich_text(self):
        source = self.make_word()
        directory = self.root / 'unpacked'
        self.run_tool('docx/ooxml/scripts/unpack.py', source, directory)
        code = '''import sys
from scripts.document import Document
document = Document(sys.argv[1], author='Restore test', initials='RT')
run = document['word/document.xml'].get_node(tag='w:r', contains='rich text')
document.add_comment(start=run, end=run, text='Restored comment')
document.save()
'''
        result = subprocess.run([sys.executable, '-c', code, str(directory)], cwd=SKILLS / 'docx',
                                text=True, capture_output=True, timeout=45)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        output = self.root / 'commented.docx'
        self.run_tool('docx/ooxml/scripts/pack.py', directory, output)
        self.assertTrue(Document(output).paragraphs[0].runs[1].bold)
        with zipfile.ZipFile(output) as package:
            self.assertIn(b'Restored comment', package.read('word/comments.xml'))

    def test_html_to_editable_powerpoint(self):
        html = self.root / 'slide.html'
        html.write_text('<!DOCTYPE html><html style="width:960px;height:540px;margin:0;padding:0">'
                        '<body style="position:relative;margin:0;padding:0;width:960px;height:540px;background:#ffffff">'
                        '<h1 style="position:absolute;left:48px;top:40px;width:800px;margin:0;font-size:32px;color:#222222">'
                        'Restored HTML capability</h1></body></html>')
        output = self.root / 'html.pptx'
        code = '''const pptxgen = require('pptxgenjs');
const html2pptx = require('./scripts/html2pptx');
(async()=>{const deck=new pptxgen();deck.layout='LAYOUT_16x9';
await html2pptx(process.argv[1],deck);await deck.writeFile({fileName:process.argv[2]});
})().catch(error=>{console.error(error);process.exitCode=1;});'''
        result = subprocess.run(['node', '-e', code, str(html), str(output)], cwd=SKILLS / 'pptx',
                                text=True, capture_output=True, timeout=90)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        deck = Presentation(output)
        self.assertTrue(any('Restored HTML capability' in getattr(shape, 'text', '') for shape in deck.slides[0].shapes))

    def test_powerpoint_inventory_and_duplicate_reorder(self):
        source = self.root / 'input.pptx'
        deck = Presentation()
        for title in ['First', 'Second']:
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = title
            slide.notes_slide.notes_text_frame.text = title + ' notes'
        deck.save(source)
        inventory = self.root / 'inventory.json'
        self.run_tool('pptx/scripts/inventory.py', source, inventory)
        self.assertIn('Second', inventory.read_text())
        output = self.root / 'reordered.pptx'
        self.run_tool('pptx/scripts/rearrange.py', source, output, '1,0,1')
        restored = Presentation(output)
        self.assertEqual([slide.shapes[0].text for slide in restored.slides], ['Second', 'First', 'Second'])
        self.assertEqual(restored.slides[0].notes_slide.notes_text_frame.text, 'Second notes')

    def test_pdf_text_annotations(self):
        source = self.make_pdf()
        fields = self.root / 'fields.json'
        fields.write_text(json.dumps({'pages': [{'page_number': 1, 'image_width': 600, 'image_height': 800}],
                                     'form_fields': [{'page_number': 1, 'description': 'Name',
                                                      'entry_bounding_box': [100, 100, 300, 125],
                                                      'entry_text': {'text': 'Restored annotation'}}]}))
        output = self.root / 'annotated.pdf'
        self.run_tool('pdf/scripts/fill_pdf_form_with_annotations.py', source, fields, output)
        annotations = PdfReader(output).pages[0]['/Annots']
        self.assertEqual(annotations[0].get_object()['/Contents'], 'Restored annotation')
        self.assertEqual(len(PdfReader(source).pages[0].get('/Annots', [])), 0)

    def test_pdf_existing_form_fields(self):
        source = self.make_pdf(form=True)
        fields = self.root / 'values.json'
        fields.write_text(json.dumps([{'field_id': 'name', 'page': 1, 'value': 'Restored value'}]))
        output = self.root / 'filled.pdf'
        self.run_tool('pdf/scripts/fill_fillable_fields.py', source, fields, output)
        self.assertEqual(PdfReader(output).get_fields()['name']['/V'], 'Restored value')

    def test_pdf_page_rasterization(self):
        source = self.make_pdf()
        output = self.root / 'images'
        output.mkdir()
        self.run_tool('pdf/scripts/convert_pdf_to_images.py', source, output)
        with Image.open(output / 'page_1.png') as picture:
            self.assertGreater(min(picture.size), 0)
            self.assertLessEqual(max(picture.size), 1000)

    def test_excel_recalculation(self):
        source = self.root / 'formula.xlsx'
        workbook = Workbook()
        workbook.active.append([24, 46, '=SUM(A1:B1)'])
        workbook.save(source)
        result = json.loads(self.run_tool('xlsx/recalc.py', source))
        self.assertEqual(result.get('status'), 'success', result)
        recalculated = load_workbook(source, data_only=True)
        self.assertEqual(recalculated.active['C1'].value, 70)
        recalculated.close()


if __name__ == '__main__':
    unittest.main(verbosity=2)
