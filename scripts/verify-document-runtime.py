"""Generate and independently reopen the four supported document formats."""
import argparse
import json
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader
from reportlab.pdfgen import canvas


def verify(directory):
    output = Path(directory)
    output.mkdir(parents=True, exist_ok=True)
    title = "StarPath document runtime"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = title
    deck.save(output / "runtime-check.pptx")
    assert Presentation(output / "runtime-check.pptx").slides[0].shapes[0].text == title
    document = Document()
    document.add_heading(title)
    document.save(output / "runtime-check.docx")
    assert Document(output / "runtime-check.docx").paragraphs[0].text == title
    workbook = Workbook()
    workbook.active.append(["Item", "Value"])
    workbook.active.append(["Example", 70])
    workbook.save(output / "runtime-check.xlsx")
    loaded = load_workbook(output / "runtime-check.xlsx", read_only=True)
    assert loaded.active["B2"].value == 70
    loaded.close()
    pdf = canvas.Canvas(str(output / "runtime-check.pdf"))
    pdf.drawString(72, 720, title)
    pdf.save()
    reader = PdfReader(output / "runtime-check.pdf")
    assert len(reader.pages) == 1 and title in reader.pages[0].extract_text()
    print(json.dumps({"verified": ["pptx", "docx", "xlsx", "pdf"]}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output")
    args = parser.parse_args()
    if args.output:
        verify(args.output)
    else:
        with TemporaryDirectory(prefix="starpath-document-check-") as directory:
            verify(directory)
