"""StarPath spreadsheet inspection and optional local recalculation. MIT."""
import argparse
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
from tempfile import TemporaryDirectory

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'local-tools' / 'scripts'))
from starpath_document_io import save_new
from openpyxl import load_workbook


def inspect(source):
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        return {'sheets': [{
            'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column,
            'formulas': [cell.coordinate for row in sheet for cell in row if cell.data_type == 'f'],
            'errors': [{'cell': cell.coordinate, 'value': cell.value} for row in sheet for cell in row if cell.data_type == 'e'],
        } for sheet in workbook]}
    finally:
        workbook.close()


def recalc(source, output):
    source = Path(source).resolve()
    output = Path(output).resolve()
    if source.suffix.lower() != '.xlsx' or output.suffix.lower() != '.xlsx':
        raise ValueError('Recalculation helper accepts XLSX only; macros require application-specific handling')
    if output.exists():
        raise FileExistsError(output)
    command = os.environ.get('STARPATH_LIBREOFFICE') or shutil.which('libreoffice') or shutil.which('soffice')
    if not command:
        raise RuntimeError('Local LibreOffice is unavailable; formulas were not recalculated')
    with TemporaryDirectory(prefix='starpath-recalc-') as temporary:
        root = Path(temporary)
        inputs = root / 'input'
        inputs.mkdir()
        shutil.copy2(source, inputs / source.name)
        destination = root / 'output'
        destination.mkdir()
        result = subprocess.run([
            command, f'-env:UserInstallation={(root / "profile").as_uri()}',
            '--headless', '--convert-to', 'xlsx', '--outdir', str(destination), str(inputs / source.name),
        ], capture_output=True, timeout=120, check=False)
        generated = destination / source.name
        if result.returncode != 0 or not generated.is_file():
            raise RuntimeError('LibreOffice did not produce a recalculated workbook')
        workbook = load_workbook(generated, read_only=True, data_only=True)
        workbook.close()
        save_new(output, lambda stream: stream.write(generated.read_bytes()))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    commands.add_parser('inspect').add_argument('input')
    command = commands.add_parser('recalc')
    command.add_argument('input')
    command.add_argument('output')
    args = parser.parse_args()
    if args.command == 'inspect':
        print(json.dumps(inspect(args.input), ensure_ascii=False, indent=2))
    else:
        recalc(args.input, args.output)


if __name__ == '__main__':
    main()
